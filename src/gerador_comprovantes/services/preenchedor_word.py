"""Preenchimento de placeholders em um documento Word (.docx)."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.document import Document as DocumentType


class PreenchedorWord:
    """
    Substitui placeholders (ex.: {{Nome}}) por valores reais em um
    documento .docx, cobrindo parágrafos, tabelas, cabeçalhos e rodapés.
    """

    def preencher_e_salvar(
        self, caminho_modelo: Path, mapa_substituicao: dict[str, str], caminho_destino: Path
    ) -> None:
        documento = Document(caminho_modelo)
        self._substituir_no_documento(documento, mapa_substituicao)
        documento.save(caminho_destino)

    def _substituir_no_documento(self, documento: DocumentType, mapa: dict[str, str]) -> None:
        self._substituir_em_paragrafos(documento.paragraphs, mapa)

        for tabela in documento.tables:
            for linha in tabela.rows:
                for celula in linha.cells:
                    self._substituir_em_paragrafos(celula.paragraphs, mapa)

        for secao in documento.sections:
            self._substituir_em_paragrafos(secao.header.paragraphs, mapa)
            self._substituir_em_paragrafos(secao.footer.paragraphs, mapa)

    @staticmethod
    def _substituir_em_paragrafos(paragrafos, mapa: dict[str, str]) -> None:
        for paragrafo in paragrafos:
            for run in paragrafo.runs:
                for chave, valor in mapa.items():
                    if chave in run.text:
                        run.text = run.text.replace(chave, valor)
